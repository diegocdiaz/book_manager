"""Servicios del sistema Book Manager.

Aca va la logica de negocio: validaciones antes de guardar, controles para
no borrar datos que se estan usando y los calculos de precios segun la
cotizacion del dolar.
"""

from __future__ import annotations

import csv
import datetime
import difflib
import os
import re
import unicodedata
from pathlib import Path
from typing import Any, Dict, List, Optional

import requests
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from book_manager.entities.entities import (
    CotizacionDolar,
    Editorial,
    EntidadBase,
    Genero,
    Libro,
    Moneda,
    Precio,
    PrecioWeb,
    Stock,
    TipoCotizacion,
)
from book_manager.repositories.repositories import (
    RepositorioCotizacionDolar,
    RepositorioDB,
    RepositorioEditorial,
    RepositorioGenero,
    RepositorioLibro,
    RepositorioMoneda,
    RepositorioPrecio,
    RepositorioPrecioWeb,
    RepositorioStock,
    RepositorioTipoCotizacion,
)


class ServicioBase:
    """Metodos que comparten los servicios de las entidades con id."""

    nombre_entidad = 'registro'

    def __init__(self, repositorio: RepositorioDB) -> None:
        self._repositorio = repositorio

    def listar(self) -> List[EntidadBase]:
        """Devuelve todos los registros."""
        return self._repositorio.leer_todos()

    def obtener(self, id_: int) -> EntidadBase:
        """Busca por id. Si no existe lanza ValueError."""
        entidad = self._repositorio.leer_por_id(id_)
        if entidad is None:
            raise ValueError(f'No existe {self.nombre_entidad} con id {id_}.')
        return entidad


class ServicioGenero(ServicioBase):
    """Logica de los generos."""

    nombre_entidad = 'genero'

    def __init__(self, repositorio: RepositorioGenero,
                 repo_libro: RepositorioLibro) -> None:
        super().__init__(repositorio)
        self._repo_libro = repo_libro

    def _existe_nombre(self, nombre: str, id_actual: int = 0) -> bool:
        for genero in self.listar():
            if (genero.nombre.lower() == nombre.strip().lower()
                    and genero.id != id_actual):
                return True
        return False

    def crear(self, nombre: str) -> Genero:
        """Da de alta un genero nuevo."""
        if self._existe_nombre(nombre):
            raise ValueError(f'Ya existe el genero {nombre}.')
        genero = Genero(self._repositorio.siguiente_id(), nombre)
        return self._repositorio.crear(genero)

    def actualizar(self, id_: int, nombre: str) -> Genero:
        """Cambia el nombre de un genero."""
        genero = self.obtener(id_)
        if self._existe_nombre(nombre, id_):
            raise ValueError(f'Ya existe el genero {nombre}.')
        genero.nombre = nombre
        return self._repositorio.actualizar(genero)

    def eliminar(self, id_: int) -> bool:
        """Borra un genero si ningun libro lo usa."""
        self.obtener(id_)
        for libro in self._repo_libro.leer_todos():
            if libro.genero.id == id_:
                raise ValueError('No se puede borrar, hay libros con ese '
                                 'genero.')
        return self._repositorio.eliminar(id_)


class ServicioEditorial(ServicioBase):
    """Logica de las editoriales."""

    nombre_entidad = 'editorial'

    def __init__(self, repositorio: RepositorioEditorial,
                 repo_libro: RepositorioLibro) -> None:
        super().__init__(repositorio)
        self._repo_libro = repo_libro

    def _existe_nombre(self, nombre: str, id_actual: int = 0) -> bool:
        for editorial in self.listar():
            if (editorial.nombre.lower() == nombre.strip().lower()
                    and editorial.id != id_actual):
                return True
        return False

    def crear(self, nombre: str, contacto: Optional[str] = None) -> Editorial:
        """Da de alta una editorial nueva."""
        if self._existe_nombre(nombre):
            raise ValueError(f'Ya existe la editorial {nombre}.')
        editorial = Editorial(self._repositorio.siguiente_id(), nombre,
                              contacto or None)
        return self._repositorio.crear(editorial)

    def actualizar(self, id_: int, nombre: str,
                   contacto: Optional[str] = None) -> Editorial:
        """Modifica nombre y contacto de una editorial."""
        editorial = self.obtener(id_)
        if self._existe_nombre(nombre, id_):
            raise ValueError(f'Ya existe la editorial {nombre}.')
        editorial.nombre = nombre
        editorial.contacto = contacto or None
        return self._repositorio.actualizar(editorial)

    def eliminar(self, id_: int) -> bool:
        """Borra una editorial si no tiene libros."""
        self.obtener(id_)
        for libro in self._repo_libro.leer_todos():
            if libro.editorial.id == id_:
                raise ValueError('No se puede borrar, la editorial tiene '
                                 'libros cargados.')
        return self._repositorio.eliminar(id_)


class ServicioMoneda(ServicioBase):
    """Logica de las monedas."""

    nombre_entidad = 'moneda'

    def __init__(self, repositorio: RepositorioMoneda,
                 repo_precio: RepositorioPrecio) -> None:
        super().__init__(repositorio)
        self._repo_precio = repo_precio

    def _existe_codigo(self, codigo: str, id_actual: int = 0) -> bool:
        for moneda in self.listar():
            if (moneda.codigo == codigo.strip().upper()
                    and moneda.id != id_actual):
                return True
        return False

    def crear(self, codigo: str, nombre: str) -> Moneda:
        """Da de alta una moneda nueva."""
        if self._existe_codigo(codigo):
            raise ValueError(f'Ya existe la moneda {codigo.upper()}.')
        moneda = Moneda(self._repositorio.siguiente_id(), codigo, nombre)
        return self._repositorio.crear(moneda)

    def actualizar(self, id_: int, codigo: str, nombre: str) -> Moneda:
        """Modifica codigo y nombre de una moneda."""
        moneda = self.obtener(id_)
        if self._existe_codigo(codigo, id_):
            raise ValueError(f'Ya existe la moneda {codigo.upper()}.')
        # primero pruebo con una moneda auxiliar para que si algun dato
        # esta mal no quede la original cambiada a medias
        Moneda(id_, codigo, nombre)
        moneda.codigo = codigo
        moneda.nombre = nombre
        return self._repositorio.actualizar(moneda)

    def eliminar(self, id_: int) -> bool:
        """Borra una moneda si no hay precios en esa moneda."""
        self.obtener(id_)
        for precio in self._repo_precio.leer_todos():
            if precio.moneda.id == id_:
                raise ValueError('No se puede borrar, hay precios en esa '
                                 'moneda.')
        return self._repositorio.eliminar(id_)


class ServicioTipoCotizacion(ServicioBase):
    """Logica de los tipos de dolar (Oficial, Blue, MEP, etc)."""

    nombre_entidad = 'tipo de cotizacion'

    def __init__(self, repositorio: RepositorioTipoCotizacion,
                 repo_cotizacion: RepositorioCotizacionDolar) -> None:
        super().__init__(repositorio)
        self._repo_cotizacion = repo_cotizacion

    def _existe_nombre(self, nombre: str, id_actual: int = 0) -> bool:
        for tipo in self.listar():
            if (tipo.nombre.lower() == nombre.strip().lower()
                    and tipo.id != id_actual):
                return True
        return False

    def crear(self, nombre: str) -> TipoCotizacion:
        """Da de alta un tipo de cotizacion."""
        if self._existe_nombre(nombre):
            raise ValueError(f'Ya existe el tipo {nombre}.')
        tipo = TipoCotizacion(self._repositorio.siguiente_id(), nombre)
        return self._repositorio.crear(tipo)

    def actualizar(self, id_: int, nombre: str) -> TipoCotizacion:
        """Cambia el nombre de un tipo de cotizacion."""
        tipo = self.obtener(id_)
        if self._existe_nombre(nombre, id_):
            raise ValueError(f'Ya existe el tipo {nombre}.')
        tipo.nombre = nombre
        return self._repositorio.actualizar(tipo)

    def eliminar(self, id_: int) -> bool:
        """Borra un tipo si no tiene cotizaciones cargadas."""
        self.obtener(id_)
        if self._repo_cotizacion.leer_historico_por_tipo(id_):
            raise ValueError('No se puede borrar, ese tipo tiene '
                             'cotizaciones cargadas.')
        return self._repositorio.eliminar(id_)


class ServicioLibro(ServicioBase):
    """Logica de los libros."""

    nombre_entidad = 'libro'

    def __init__(self, repositorio: RepositorioLibro,
                 servicio_editorial: ServicioEditorial,
                 servicio_genero: ServicioGenero,
                 repo_precio: RepositorioPrecio,
                 repo_stock: RepositorioStock) -> None:
        super().__init__(repositorio)
        self._servicio_editorial = servicio_editorial
        self._servicio_genero = servicio_genero
        self._repo_precio = repo_precio
        self._repo_stock = repo_stock

    def limpiar_isbn(self, isbn: str) -> str:
        """Saca guiones y espacios y controla que tenga 10 o 13 digitos.

        El ISBN-10 puede terminar en X.
        """
        isbn = isbn.replace('-', '').replace(' ', '').upper()
        if len(isbn) == 13 and isbn.isdigit():
            return isbn
        if len(isbn) == 10 and isbn[:9].isdigit() and (
                isbn[9].isdigit() or isbn[9] == 'X'):
            return isbn
        raise ValueError('El ISBN tiene que tener 10 o 13 digitos.')

    def crear(self, isbn: str, titulo: str, autor: str,
              editorial_id: int, genero_id: int) -> Libro:
        """Da de alta un libro. El ISBN no se puede repetir."""
        isbn = self.limpiar_isbn(isbn)
        if self._repositorio.leer_por_isbn(isbn) is not None:
            raise ValueError(f'Ya existe un libro con ISBN {isbn}.')
        editorial = self._servicio_editorial.obtener(editorial_id)
        genero = self._servicio_genero.obtener(genero_id)
        libro = Libro(self._repositorio.siguiente_id(), isbn, titulo, autor,
                      editorial, genero)
        return self._repositorio.crear(libro)

    def actualizar(self, id_: int, isbn: str, titulo: str, autor: str,
                   editorial_id: int, genero_id: int) -> Libro:
        """Modifica los datos de un libro."""
        libro = self.obtener(id_)
        isbn = self.limpiar_isbn(isbn)
        otro = self._repositorio.leer_por_isbn(isbn)
        if otro is not None and otro.id != id_:
            raise ValueError(f'Ya existe un libro con ISBN {isbn}.')
        editorial = self._servicio_editorial.obtener(editorial_id)
        genero = self._servicio_genero.obtener(genero_id)
        # pruebo los datos con un libro auxiliar antes de tocar el original.
        # No reemplazo el objeto porque Precio y Stock apuntan a este mismo
        Libro(id_, isbn, titulo, autor, editorial, genero)
        libro.isbn = isbn
        libro.titulo = titulo
        libro.autor = autor
        libro.editorial = editorial
        libro.genero = genero
        return self._repositorio.actualizar(libro)

    def eliminar(self, id_: int) -> bool:
        """Borra un libro y tambien sus precios y su stock."""
        self.obtener(id_)
        for precio in self._repo_precio.leer_por_libro(id_):
            self._repo_precio.eliminar(precio.id)
        self._repo_stock.eliminar(id_)
        return self._repositorio.eliminar(id_)


class ServicioPrecio(ServicioBase):
    """Logica de los precios. Cada libro tiene un solo precio por moneda."""

    nombre_entidad = 'precio'

    def __init__(self, repositorio: RepositorioPrecio,
                 servicio_libro: ServicioLibro,
                 servicio_moneda: ServicioMoneda) -> None:
        super().__init__(repositorio)
        self._servicio_libro = servicio_libro
        self._servicio_moneda = servicio_moneda

    def buscar_precio(self, libro_id: int, codigo: str) -> Optional[Precio]:
        """Devuelve el precio del libro en esa moneda o None si no tiene."""
        for precio in self._repositorio.leer_por_libro(libro_id):
            if precio.moneda.codigo == codigo.upper():
                return precio
        return None

    def precios_por_libro(self) -> Dict[int, Dict[str, float]]:
        """Agrupa todos los precios por libro y por codigo de moneda.

        Returns:
            Dict[int, Dict[str, float]]: Para cada id de libro, el valor de
                su precio en cada moneda (ej: {2: {'USD': 26.0}}).
        """
        precios: Dict[int, Dict[str, float]] = {}
        for precio in self.listar():
            precios.setdefault(precio.libro.id, {})[precio.moneda.codigo] = (
                precio.valor)
        return precios

    def crear(self, libro_id: int, moneda_id: int, valor: float) -> Precio:
        """Carga el precio de un libro en una moneda."""
        libro = self._servicio_libro.obtener(libro_id)
        moneda = self._servicio_moneda.obtener(moneda_id)
        if self.buscar_precio(libro_id, moneda.codigo) is not None:
            raise ValueError(f'{libro.titulo} ya tiene precio en '
                             f'{moneda.codigo}.')
        if valor <= 0:
            raise ValueError('El precio tiene que ser mayor a 0.')
        precio = Precio(self._repositorio.siguiente_id(), libro, moneda,
                        valor)
        return self._repositorio.crear(precio)

    def actualizar(self, id_: int, valor: float) -> Precio:
        """Cambia el valor de un precio."""
        precio = self.obtener(id_)
        if valor <= 0:
            raise ValueError('El precio tiene que ser mayor a 0.')
        precio.valor = valor
        return self._repositorio.actualizar(precio)

    def eliminar(self, id_: int) -> bool:
        """Borra un precio."""
        self.obtener(id_)
        return self._repositorio.eliminar(id_)


class ServicioStock:
    """Logica del stock. Se busca por el id del libro."""

    def __init__(self, repositorio: RepositorioStock,
                 servicio_libro: ServicioLibro) -> None:
        self._repositorio = repositorio
        self._servicio_libro = servicio_libro

    def listar(self) -> List[Stock]:
        """Devuelve todo el stock."""
        return self._repositorio.leer_todos()

    def obtener(self, libro_id: int) -> Stock:
        """Devuelve el stock de un libro."""
        stock = self._repositorio.leer_por_libro(libro_id)
        if stock is None:
            raise ValueError(f'El libro {libro_id} no tiene stock cargado.')
        return stock

    def crear(self, libro_id: int, cantidad: int) -> Stock:
        """Carga el stock inicial de un libro."""
        libro = self._servicio_libro.obtener(libro_id)
        stock = Stock(self._repositorio.siguiente_id(), libro, cantidad)
        return self._repositorio.crear(stock)

    def actualizar(self, libro_id: int, cantidad: int) -> Stock:
        """Pone la cantidad que se indica (por ejemplo despues de contar)."""
        stock = self.obtener(libro_id)
        stock.cantidad = cantidad
        return self._repositorio.actualizar(stock)

    def eliminar(self, libro_id: int) -> bool:
        """Borra el stock de un libro."""
        self.obtener(libro_id)
        return self._repositorio.eliminar(libro_id)


# Archivo .env en la raiz del repositorio (book_manager/.env)
RUTA_ENV = Path(__file__).resolve().parents[3] / '.env'

# Zona horaria de Argentina (UTC-3) para pasar la fecha de la API a dia local
ZONA_ARGENTINA = datetime.timezone(datetime.timedelta(hours=-3))

# Relacion entre el campo "casa" de la API y los tipos de cotizacion del
# sistema. Si aparece una casa nueva se usa el nombre que trae la API.
TIPOS_POR_CASA = {
    'oficial': 'Oficial',
    'blue': 'Blue',
    'bolsa': 'MEP',
    'contadoconliqui': 'CCL',
    'mayorista': 'Mayorista',
    'cripto': 'Cripto',
    'tarjeta': 'Tarjeta',
}


class ServicioCotizacionDolar:
    """Logica de las cotizaciones del dolar."""

    def __init__(self, repositorio: RepositorioCotizacionDolar,
                 servicio_tipo: ServicioTipoCotizacion) -> None:
        self._repositorio = repositorio
        self._servicio_tipo = servicio_tipo

    def listar(self) -> List[CotizacionDolar]:
        """Devuelve todas las cotizaciones."""
        return self._repositorio.leer_todos()

    def historico(self, tipo_id: int) -> List[CotizacionDolar]:
        """Cotizaciones de un tipo ordenadas por fecha."""
        self._servicio_tipo.obtener(tipo_id)
        return self._repositorio.leer_historico_por_tipo(tipo_id)

    def obtener(self, tipo_id: int, fecha: datetime.date) -> CotizacionDolar:
        """Busca la cotizacion de un tipo en una fecha."""
        cotizacion = self._repositorio.leer_por_tipo_y_fecha(tipo_id, fecha)
        if cotizacion is None:
            raise ValueError(f'No hay cotizacion para esa fecha ({fecha}).')
        return cotizacion

    def ultima(self, tipo_id: int) -> CotizacionDolar:
        """Devuelve la cotizacion mas nueva de un tipo."""
        historico = self.historico(tipo_id)
        if not historico:
            raise ValueError('Ese tipo de dolar no tiene cotizaciones.')
        return historico[-1]

    def ultimas(self) -> List[CotizacionDolar]:
        """Devuelve la cotizacion mas nueva de cada tipo que tiene datos.

        Se leen todas las cotizaciones en una sola consulta y se queda la de
        fecha mas reciente de cada tipo.

        Returns:
            List[CotizacionDolar]: Una cotizacion por tipo, ordenadas por el
                id del tipo.
        """
        ultimas: Dict[int, CotizacionDolar] = {}
        for cotizacion in self.listar():
            actual = ultimas.get(cotizacion.tipo.id)
            if actual is None or cotizacion.fecha > actual.fecha:
                ultimas[cotizacion.tipo.id] = cotizacion
        return [ultimas[tipo_id] for tipo_id in sorted(ultimas)]

    def _validar(self, fecha: datetime.date, compra: float,
                 venta: float) -> None:
        if fecha > datetime.date.today():
            raise ValueError('La fecha no puede ser futura.')
        if compra <= 0 or venta <= 0:
            raise ValueError('Compra y venta tienen que ser mayores a 0.')
        if venta < compra:
            raise ValueError('La venta no puede ser menor que la compra.')

    def crear(self, tipo_id: int, fecha: datetime.date, compra: float,
              venta: float) -> CotizacionDolar:
        """Carga una cotizacion nueva."""
        tipo = self._servicio_tipo.obtener(tipo_id)
        self._validar(fecha, compra, venta)
        cotizacion = CotizacionDolar(self._repositorio.siguiente_id(), tipo,
                                     fecha, compra, venta)
        return self._repositorio.crear(cotizacion)

    def actualizar(self, tipo_id: int, fecha: datetime.date, compra: float,
                   venta: float) -> CotizacionDolar:
        """Corrige los valores de una cotizacion."""
        cotizacion = self.obtener(tipo_id, fecha)
        self._validar(fecha, compra, venta)
        cotizacion.valor_compra = compra
        cotizacion.valor_venta = venta
        return self._repositorio.actualizar(cotizacion)

    def eliminar(self, tipo_id: int, fecha: datetime.date) -> bool:
        """Borra una cotizacion."""
        self.obtener(tipo_id, fecha)
        return self._repositorio.eliminar(tipo_id, fecha)

    def _descargar(self, timeout: int) -> List[Dict[str, Any]]:
        """Descarga las cotizaciones de la API configurada en el .env."""
        load_dotenv(RUTA_ENV)
        url = os.getenv('API_URL')
        if not url:
            raise ValueError('Falta la variable API_URL en el archivo .env.')
        try:
            respuesta = requests.get(url, timeout=timeout)
            respuesta.raise_for_status()
            return respuesta.json()
        except (requests.RequestException, ValueError) as error:
            raise ValueError(
                f'No se pudieron obtener las cotizaciones: {error}'
            ) from error

    def _tipo_para(self, dato: Dict[str, Any],
                   guardar: bool) -> TipoCotizacion:
        """Busca el tipo de cotizacion que corresponde a un dato de la API.

        Si no existe y se va a guardar, se da de alta.
        """
        nombre = TIPOS_POR_CASA.get(dato['casa'], dato['nombre'])
        for tipo in self._servicio_tipo.listar():
            if tipo.nombre.lower() == nombre.lower():
                return tipo
        if guardar:
            return self._servicio_tipo.crear(nombre)
        return TipoCotizacion(0, nombre)

    @staticmethod
    def _fecha_api(texto: str) -> datetime.date:
        """Convierte la fecha de la API (UTC) en la fecha de Argentina."""
        momento = datetime.datetime.fromisoformat(texto.replace('Z', '+00:00'))
        return momento.astimezone(ZONA_ARGENTINA).date()

    def obtener_cotizaciones(self, guardar: bool = True,
                             timeout: int = 10) -> List[CotizacionDolar]:
        """Busca, descarga y registra las cotizaciones del dolar.

        Consulta la API indicada en la variable API_URL del archivo .env.
        Por cada tipo de dolar que devuelve la API se registra la cotizacion
        del dia: si ya habia una para ese tipo y fecha se actualizan los
        valores, si no se crea. Los datos sin compra o sin venta se omiten.

        Args:
            guardar (bool): Si es False solo se descargan las cotizaciones,
                sin registrar nada en la base.
            timeout (int): Segundos maximos de espera de la API.

        Returns:
            List[CotizacionDolar]: Las cotizaciones obtenidas.

        Raises:
            ValueError: Si falta API_URL o la API no responde bien.
        """
        cotizaciones = []
        for dato in self._descargar(timeout):
            compra, venta = dato.get('compra'), dato.get('venta')
            if compra is None or venta is None:
                continue
            tipo = self._tipo_para(dato, guardar)
            fecha = self._fecha_api(dato['fechaActualizacion'])
            if not guardar:
                cotizaciones.append(
                    CotizacionDolar(0, tipo, fecha, compra, venta)
                )
            elif self._repositorio.leer_por_tipo_y_fecha(tipo.id, fecha):
                cotizaciones.append(
                    self.actualizar(tipo.id, fecha, compra, venta)
                )
            else:
                cotizaciones.append(self.crear(tipo.id, fecha, compra, venta))
        return cotizaciones


# Columnas del CSV con la lista de precios bimonetaria
ENCABEZADOS_EXPORTACION = (
    'libro_id', 'isbn', 'titulo', 'autor', 'tipo_dolar', 'fecha_cotizacion',
    'dolar_venta', 'precio_ars', 'precio_usd',
)


class ServicioCotizador:
    """Calcula los precios de los libros en pesos y dolares.

    Si el libro tiene precio en USD se pasa a pesos con el dolar venta de
    la ultima cotizacion del tipo elegido. Si solo tiene precio en ARS se
    hace la cuenta al reves.
    """

    def __init__(self, servicio_libro: ServicioLibro,
                 servicio_precio: ServicioPrecio,
                 servicio_cotizacion: ServicioCotizacionDolar) -> None:
        self._servicio_libro = servicio_libro
        self._servicio_precio = servicio_precio
        self._servicio_cotizacion = servicio_cotizacion

    @staticmethod
    def _cotizar(libro: Libro, cotizacion: CotizacionDolar,
                 precios: Dict[str, float]) -> Optional[Dict]:
        """Pasa el precio de un libro a ARS y USD con una cotizacion.

        Args:
            libro (Libro): Libro a cotizar.
            cotizacion (CotizacionDolar): Cotizacion cuyo valor de venta se
                usa para la conversion.
            precios (Dict[str, float]): Precios del libro por codigo de
                moneda.

        Returns:
            Optional[Dict]: Los datos de la cotizacion del libro, o None si
                no tiene precio en USD ni en ARS.
        """
        dolar = cotizacion.valor_venta
        if 'USD' in precios:
            usd = precios['USD']
            ars = usd * dolar
        elif 'ARS' in precios:
            ars = precios['ARS']
            usd = ars / dolar
        else:
            return None

        return {
            'libro': libro,
            'tipo': cotizacion.tipo.nombre,
            'fecha': cotizacion.fecha,
            'dolar': dolar,
            'ars': round(ars, 2),
            'usd': round(usd, 2),
        }

    def _cotizar_catalogo(self, libros: List[Libro],
                          cotizacion: CotizacionDolar,
                          precios: Dict[int, Dict[str, float]]) -> List[Dict]:
        """Cotiza con una misma cotizacion los libros que tienen precio."""
        filas = (self._cotizar(libro, cotizacion, precios.get(libro.id, {}))
                 for libro in libros)
        return [fila for fila in filas if fila is not None]

    def cotizar_libro(self, libro_id: int, tipo_id: int) -> Dict:
        """Devuelve un diccionario con el precio del libro en ARS y USD."""
        libro = self._servicio_libro.obtener(libro_id)
        cotizacion = self._servicio_cotizacion.ultima(tipo_id)

        precios = {}
        for codigo in ('USD', 'ARS'):
            precio = self._servicio_precio.buscar_precio(libro_id, codigo)
            if precio is not None:
                precios[codigo] = precio.valor

        fila = self._cotizar(libro, cotizacion, precios)
        if fila is None:
            raise ValueError(f'{libro.titulo} no tiene precio cargado.')
        return fila

    def cotizar_todos(self, tipo_id: int) -> List[Dict]:
        """Cotiza todos los libros que tienen precio."""
        cotizacion = self._servicio_cotizacion.ultima(tipo_id)
        precios = self._servicio_precio.precios_por_libro()
        return self._cotizar_catalogo(self._servicio_libro.listar(),
                                      cotizacion, precios)

    def exportar_precios_csv(self, ruta: Path) -> int:
        """Exporta a CSV los precios en ARS y USD para cada tipo de dolar.

        Se genera una fila por cada libro con precio y cada tipo de
        cotizacion que tenga datos, usando la ultima cotizacion de ese tipo.

        Args:
            ruta (Path): Archivo CSV a generar. Si la carpeta no existe se
                crea.

        Returns:
            int: Cantidad de filas exportadas.

        Raises:
            ValueError: Si no hay cotizaciones o ningun libro tiene precio.
        """
        cotizaciones = self._servicio_cotizacion.ultimas()
        if not cotizaciones:
            raise ValueError('No hay cotizaciones del dolar cargadas.')

        libros = self._servicio_libro.listar()
        precios = self._servicio_precio.precios_por_libro()
        filas = [fila for cotizacion in cotizaciones
                 for fila in self._cotizar_catalogo(libros, cotizacion,
                                                    precios)]
        if not filas:
            raise ValueError('Ningun libro tiene precio cargado.')

        ruta.parent.mkdir(parents=True, exist_ok=True)
        with ruta.open('w', newline='', encoding='utf-8') as archivo:
            escritor = csv.writer(archivo)
            escritor.writerow(ENCABEZADOS_EXPORTACION)
            escritor.writerows(
                (fila['libro'].id, fila['libro'].isbn, fila['libro'].titulo,
                 fila['libro'].autor, fila['tipo'], fila['fecha'].isoformat(),
                 fila['dolar'], fila['ars'], fila['usd'])
                for fila in filas
            )
        return len(filas)

    def comparar_competencia(self, libro_id: int, tipo_id: int,
                             precio_competencia: float,
                             competidor: str = 'Cuspide') -> Dict:
        """Compara nuestro precio en pesos con el de otra libreria."""
        if precio_competencia <= 0:
            raise ValueError('El precio de la competencia tiene que ser '
                             'mayor a 0.')
        nuestro = self.cotizar_libro(libro_id, tipo_id)['ars']
        diferencia = nuestro - precio_competencia
        return {
            'libro': self._servicio_libro.obtener(libro_id),
            'competidor': competidor,
            'nuestro': nuestro,
            'competencia': precio_competencia,
            'diferencia': round(diferencia, 2),
            'porcentaje': round(diferencia / precio_competencia * 100, 2),
            'mas_barato': nuestro < precio_competencia,
        }


# Columnas del CSV de alertas de precios
ENCABEZADOS_ALERTAS = (
    'libro_id', 'isbn', 'titulo', 'titulo_web', 'precio_interno',
    'precio_web', 'diferencia', 'porcentaje', 'alerta', 'fecha_extraccion',
    'url',
)

# Parecido minimo (de 0 a 1) entre dos titulos para tomarlos como el mismo
# libro
SIMILITUD_MINIMA = 0.8

# Columnas del reporte en Excel (en este orden)
ENCABEZADOS_EXCEL = ['Libro', 'Precio interno', 'Precio web', 'Diferencia',
                     'Fecha de extracción']
# Carpeta donde se guardan los reportes generados (CSV de alertas, Excel)
CARPETA_REPORTES = Path(__file__).resolve().parents[1] / 'reportes'


class ServicioComparacionPrecios:
    """Compara nuestros precios con los que el scraper saco de Cuspide.

    El scraper guarda hasta 10 resultados por cada libro buscado. Para
    cada libro se elige, entre los resultados de la ultima ejecucion, el
    que tiene el titulo mas parecido al nuestro (o el mismo ISBN) y se
    calcula la diferencia de precio en pesos.
    """

    def __init__(self, repositorio: RepositorioPrecioWeb,
                 servicio_libro: ServicioLibro,
                 cotizador: ServicioCotizador,
                 servicio_tipo: ServicioTipoCotizacion) -> None:
        self._repositorio = repositorio
        self._servicio_libro = servicio_libro
        self._cotizador = cotizador
        self._servicio_tipo = servicio_tipo

    @staticmethod
    def _normalizar(texto: str) -> str:
        """Pasa un titulo a minusculas, sin tildes ni signos."""
        sin_tildes = unicodedata.normalize('NFKD', texto)
        sin_tildes = ''.join(c for c in sin_tildes
                             if not unicodedata.combining(c))
        return re.sub(r'[^a-z0-9]+', ' ', sin_tildes.lower()).strip()

    @classmethod
    def similitud(cls, titulo: str, titulo_web: str) -> float:
        """Dice cuanto se parecen dos titulos, de 0 (nada) a 1 (iguales).

        El buscador del sitio no entiende los subtitulos, por eso tambien
        se compara solo la parte anterior a los dos puntos y se toma el
        mejor resultado.
        """
        propio = cls._normalizar(titulo)
        propio_corto = cls._normalizar(titulo.split(':')[0])
        web = cls._normalizar(titulo_web)
        web_corto = cls._normalizar(titulo_web.split(':')[0])
        pares = ((propio, web), (propio_corto, web), (propio_corto, web_corto))
        return max(difflib.SequenceMatcher(None, a, b).ratio()
                   for a, b in pares)

    def _elegir_resultado(self, libro: Libro,
                          resultados: List[PrecioWeb]) -> Optional[PrecioWeb]:
        """Devuelve el resultado del sitio que corresponde al libro."""
        mejor, mejor_puntaje = None, 0.0
        for resultado in resultados:
            if resultado.libro.id != libro.id:
                continue
            if self._mismo_isbn(libro.isbn, resultado.isbn):
                # el ISBN identifica la edicion: gana siempre sobre el
                # titulo (puede haber varias ediciones con el mismo titulo)
                puntaje = 2.0
            else:
                puntaje = self.similitud(libro.titulo, resultado.titulo)
            if puntaje > mejor_puntaje:
                mejor, mejor_puntaje = resultado, puntaje
        return mejor if mejor_puntaje >= SIMILITUD_MINIMA else None

    @staticmethod
    def _mismo_isbn(isbn: Optional[str], isbn_web: Optional[str]) -> bool:
        """Dice si dos ISBN son iguales, sin tener en cuenta los guiones."""
        if not isbn or not isbn_web:
            return False
        return re.sub(r'\D', '', isbn) == re.sub(r'\D', '', isbn_web)

    def _id_tipo_dolar(self, tipo_id: Optional[int]) -> int:
        """Devuelve el tipo de dolar a usar (por defecto, el oficial)."""
        if tipo_id is not None:
            return tipo_id
        tipos = self._servicio_tipo.listar()
        if not tipos:
            raise ValueError('No hay tipos de cotizacion cargados.')
        for tipo in tipos:
            if tipo.nombre.lower() == 'oficial':
                return tipo.id
        return tipos[0].id

    def comparar(self, tipo_id: Optional[int] = None) -> List[Dict]:
        """Compara el precio de cada libro con el de la ultima extraccion.

        Args:
            tipo_id (Optional[int]): Tipo de dolar con el que se pasan a
                pesos los libros que tienen precio solo en USD. Si no se
                indica se usa el oficial.

        Returns:
            List[Dict]: Una fila por cada libro con precio. Si no se
                encontro el libro en el sitio, precio_web, diferencia y
                porcentaje quedan en None.

        Raises:
            ValueError: Si el scraper todavia no se ejecuto.
        """
        resultados = self._repositorio.leer_ultima_extraccion()
        if not resultados:
            raise ValueError('Todavia no se ejecuto el scraper: no hay '
                             'precios de la web para comparar.')
        fecha = resultados[0].fecha_extraccion
        internos = {fila['libro'].id: fila['ars'] for fila in
                    self._cotizador.cotizar_todos(
                        self._id_tipo_dolar(tipo_id))}

        filas = []
        for libro in self._servicio_libro.listar():
            interno = internos.get(libro.id)
            if interno is None:
                continue
            web = self._elegir_resultado(libro, resultados)
            fila = {'libro': libro, 'precio_interno': interno,
                    'titulo_web': None, 'url': None, 'precio_web': None,
                    'diferencia': None, 'porcentaje': None,
                    'fecha_extraccion': fecha}
            if web is not None:
                diferencia = round(interno - web.precio, 2)
                fila.update(
                    titulo_web=web.titulo, url=web.url,
                    precio_web=web.precio, diferencia=diferencia,
                    porcentaje=(round(diferencia / web.precio * 100, 2)
                                if web.precio else None))
            filas.append(fila)
        return filas

    def generar_alertas(self, minimo: float,
                        tipo_id: Optional[int] = None) -> List[Dict]:
        """Devuelve los libros cuya diferencia de precio llega al minimo.

        Args:
            minimo (float): Diferencia minima en pesos (sin importar el
                signo) para que un libro genere una alerta.
            tipo_id (Optional[int]): Tipo de dolar para pasar a pesos.

        Returns:
            List[Dict]: Las filas de comparar() que son alerta, con la
                clave 'alerta' que dice si somos mas caros o mas baratos.

        Raises:
            ValueError: Si el minimo es negativo o el scraper no se ejecuto.
        """
        if minimo < 0:
            raise ValueError('La diferencia minima no puede ser negativa.')
        alertas = []
        for fila in self.comparar(tipo_id):
            if fila['diferencia'] is None or abs(fila['diferencia']) < minimo:
                continue
            if fila['diferencia'] > 0:
                fila['alerta'] = 'Somos mas caros'
            elif fila['diferencia'] < 0:
                fila['alerta'] = 'Somos mas baratos'
            else:
                fila['alerta'] = 'Mismo precio'
            alertas.append(fila)
        return alertas

    @staticmethod
    def ruta_alertas_por_defecto() -> Path:
        """Devuelve la ruta del CSV de alertas, con la fecha y hora."""
        momento = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
        return CARPETA_REPORTES / f'alertas_{momento}.csv'

    def exportar_alertas_csv(self, minimo: float, ruta: Path,
                             tipo_id: Optional[int] = None) -> int:
        """Guarda en un CSV las alertas de precio.

        El archivo se genera siempre, aunque no haya ninguna alerta (en ese
        caso solo tiene los encabezados).

        Args:
            minimo (float): Diferencia minima en pesos para alertar.
            ruta (Path): Archivo CSV a generar. Si la carpeta no existe se
                crea.
            tipo_id (Optional[int]): Tipo de dolar para pasar a pesos.

        Returns:
            int: Cantidad de alertas escritas.
        """
        alertas = self.generar_alertas(minimo, tipo_id)
        ruta.parent.mkdir(parents=True, exist_ok=True)
        # utf-8-sig agrega una marca al inicio para que Excel muestre bien
        # las tildes al abrir el archivo
        with ruta.open('w', newline='', encoding='utf-8-sig') as archivo:
            escritor = csv.writer(archivo)
            escritor.writerow(ENCABEZADOS_ALERTAS)
            escritor.writerows(
                (f['libro'].id, f['libro'].isbn, f['libro'].titulo,
                 f['titulo_web'], f['precio_interno'], f['precio_web'],
                 f['diferencia'], f['porcentaje'], f['alerta'],
                 f['fecha_extraccion'].isoformat(timespec='seconds'),
                 f['url'])
                for f in alertas
            )
        return len(alertas)

    @staticmethod
    def ruta_excel_por_defecto() -> Path:
        """Devuelve la ruta del reporte en Excel, con la fecha y hora."""
        momento = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
        return CARPETA_REPORTES / f'comparacion_precios_{momento}.xlsx'

    def exportar_excel(self, ruta: Path,
                       tipo_id: Optional[int] = None) -> int:
        """Guarda en un Excel la comparacion de precios de todos los libros.

        Columnas: Libro, Precio interno, Precio web, Diferencia y Fecha de
        extraccion. Los libros que no se encontraron en el sitio quedan con
        el precio web y la diferencia vacios.

        Args:
            ruta (Path): Archivo .xlsx a generar. Si la carpeta no existe se
                crea.
            tipo_id (Optional[int]): Tipo de dolar para pasar a pesos.

        Returns:
            int: Cantidad de libros escritos en el reporte.

        Raises:
            ValueError: Si el scraper todavia no se ejecuto.
        """
        filas = self.comparar(tipo_id)
        libro_excel = Workbook()
        hoja = libro_excel.active
        hoja.title = 'Comparación de precios'

        # titulo del reporte (filas 1 y 2), una fila en blanco y los
        # encabezados de la tabla en la fila 4
        hoja.merge_cells('A1:E1')
        hoja['A1'] = 'Comparación de precios: Book Manager vs. Cúspide'
        hoja['A1'].font = Font(bold=True, size=16, color='1F4E78')
        hoja['A1'].alignment = Alignment(horizontal='center')
        hoja.merge_cells('A2:E2')
        hoja['A2'] = ('Precios de la extracción del '
                      + filas[0]['fecha_extraccion'].strftime('%d/%m/%Y %H:%M')
                      + ' (diferencia en pesos: precio interno - precio web)')
        hoja['A2'].font = Font(italic=True, color='595959')
        hoja['A2'].alignment = Alignment(horizontal='center')

        fila_encabezado = 4
        hoja.append([])
        hoja.append(ENCABEZADOS_EXCEL)
        for celda in hoja[fila_encabezado]:
            celda.font = Font(bold=True, color='FFFFFF')
            celda.fill = PatternFill('solid', fgColor='1F4E78')
            celda.alignment = Alignment(horizontal='center')

        for f in filas:
            hoja.append([f['libro'].titulo, f['precio_interno'],
                         f['precio_web'], f['diferencia'],
                         f['fecha_extraccion']])
        for fila in hoja.iter_rows(min_row=fila_encabezado + 1):
            for celda in fila[1:4]:
                celda.number_format = '#,##0.00'
            fila[4].number_format = 'dd/mm/yyyy hh:mm'

        for numero, ancho in enumerate((50, 16, 16, 16, 22), start=1):
            hoja.column_dimensions[get_column_letter(numero)].width = ancho
        hoja.freeze_panes = f'A{fila_encabezado + 1}'
        hoja.auto_filter.ref = (f'A{fila_encabezado}:E'
                                f'{fila_encabezado + len(filas)}')

        # para que al imprimir entren todas las columnas en una hoja
        hoja.page_setup.orientation = 'landscape'
        hoja.sheet_properties.pageSetUpPr.fitToPage = True
        hoja.page_setup.fitToHeight = 0

        ruta.parent.mkdir(parents=True, exist_ok=True)
        libro_excel.save(ruta)
        return len(filas)


class ServiciosLibreria:
    """Agrupa todos los servicios en un solo objeto.

    Asi la consola y la precarga reciben un unico parametro. Los servicios
    llegan ya armados: la creacion de repositorios y servicios se hace en
    main.py.
    """

    def __init__(self, generos: ServicioGenero,
                 editoriales: ServicioEditorial,
                 monedas: ServicioMoneda,
                 tipos_cotizacion: ServicioTipoCotizacion,
                 libros: ServicioLibro,
                 precios: ServicioPrecio,
                 stock: ServicioStock,
                 cotizaciones: ServicioCotizacionDolar,
                 cotizador: ServicioCotizador,
                 comparador: ServicioComparacionPrecios) -> None:
        self.generos = generos
        self.editoriales = editoriales
        self.monedas = monedas
        self.tipos_cotizacion = tipos_cotizacion
        self.libros = libros
        self.precios = precios
        self.stock = stock
        self.cotizaciones = cotizaciones
        self.cotizador = cotizador
        self.comparador = comparador
